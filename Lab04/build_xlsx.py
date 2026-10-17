#!/usr/bin/env python3
"""Gera resultados_openmp.xlsx a partir dos CSVs em results/.

Abas (uma por versão de código): "Ex2 Critical", "Ex2 Reduction",
"Ex3 Sequencial", "Ex3 Paralela". Médias, speedup e eficiência são fórmulas
do Excel; os gráficos de speedup são criados automaticamente.
Uso: pip install openpyxl && python3 build_xlsx.py
"""
import csv, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

RES = "results"
THREADS = [1, 2, 3, 4, 5, 6]
FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = Font(name=FONT, color="0000FF")      # dados medidos (entrada)
CALC = Font(name=FONT, color="000000")       # fórmulas
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def read_csv(name):
    """-> {threads: [tempos]} e {threads: N}"""
    times, sizes = defaultdict(list), {}
    path = os.path.join(RES, name)
    if os.path.exists(path):
        for row in csv.reader(open(path)):
            if len(row) == 5:
                _, t, n, _, tempo = row
                times[int(t)].append(float(tempo)); sizes[int(t)] = int(n)
    return times, sizes


def header(ws, row, labels):
    for c, text in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF"); cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def title(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, bold=True, size=12)


def table(ws, top, times, sizes, nruns, mode, base_n=None):
    """Tabela: Threads | N | Exec 1..nruns | Média | Speedup | Eficiência | Ideal.
    mode: 'forte' (speedup = T1/Tp) ou 'fraca' (speedup escalado = p*T1/Tp,
    eficiência = T1/Tp). Retorna as linhas (primeira, última) de dados."""
    runs = [f"Exec {i}" for i in range(1, nruns + 1)]
    mc = 3 + nruns                                   # coluna da média
    sp_name = "Speedup" if mode == "forte" else "Speedup escalado"
    header(ws, top, ["Threads", "N"] + runs + ["Média (s)", sp_name, "Eficiência", "Ideal"])
    first = top + 1
    for k, p in enumerate(THREADS):
        r = first + k
        ws.cell(row=r, column=1, value=p).font = INPUT
        n = sizes.get(p)
        ws.cell(row=r, column=2, value=n if n else None).font = INPUT
        for i in range(nruns):
            v = times[p][i] if i < len(times.get(p, [])) else None
            c = ws.cell(row=r, column=3 + i, value=v); c.font = INPUT; c.number_format = "0.000000"
        rng = f"{L(3)}{r}:{L(2 + nruns)}{r}"
        ws.cell(row=r, column=mc, value=f'=IFERROR(AVERAGE({rng}),"")')
        m, m1 = f"{L(mc)}{r}", f"{L(mc)}${first}"
        if mode == "forte":
            sp, ef = f'=IFERROR({m1}/{m},"")', f'=IFERROR({L(mc+1)}{r}/$A{r},"")'
        else:
            sp, ef = f'=IFERROR($A{r}*{m1}/{m},"")', f'=IFERROR({m1}/{m},"")'
        ws.cell(row=r, column=mc + 1, value=sp)
        ws.cell(row=r, column=mc + 2, value=ef)
        ws.cell(row=r, column=mc + 3, value=f"=A{r}")
        for c in range(1, mc + 4):
            cell = ws.cell(row=r, column=c); cell.border = BOX
            if c >= mc and c != mc + 3: cell.font = CALC
            if c in (mc,): cell.number_format = "0.000000"
            if c in (mc + 1, mc + 2): cell.number_format = "0.00"
            if c == mc + 3: cell.font = CALC
    return first, first + len(THREADS) - 1, mc


def chart(ws, ttl, first, last, mc, anchor, ytitle):
    ch = LineChart(); ch.title = ttl; ch.height = 8.5; ch.width = 16
    ch.x_axis.title = "Threads"; ch.y_axis.title = ytitle
    ch.x_axis.delete = False; ch.y_axis.delete = False
    for col in (mc + 1, mc + 3):                     # Speedup e Ideal
        ch.add_data(Reference(ws, min_col=col, min_row=first - 1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    for s in ch.series: s.marker.symbol = "circle"
    ws.add_chart(ch, anchor)


def ex2_sheet(wb, name, version):
    ws = wb.create_sheet(name)
    ws["A1"] = f"Exercício 2 – Soma de vetor (versão {version})"; ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Azul = dados medidos (tempos em segundos, de results/*.csv); preto = fórmulas. Média de 10 execuções."
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    t, s = read_csv(f"ex2_{version}_forte.csv")
    title(ws, 4, "Escalabilidade FORTE (N fixo = 2^28 elementos)")
    f1, l1, mc = table(ws, 5, t, s, 10, "forte")
    chart(ws, "Speedup – escalabilidade forte", f1, l1, mc, f"{L(mc+5)}4", "Speedup")
    t, s = read_csv(f"ex2_{version}_fraca.csv")
    title(ws, 14, "Escalabilidade FRACA (N = threads × 2^27; máximo 6×2^27 ≤ 2^30)")
    f2, l2, mc = table(ws, 15, t, s, 10, "fraca")
    chart(ws, "Speedup escalado – escalabilidade fraca", f2, l2, mc, f"{L(mc+5)}22", "Speedup escalado")
    ws.column_dimensions["A"].width = 9; ws.column_dimensions["B"].width = 13
    for c in range(3, mc + 4): ws.column_dimensions[L(c)].width = 11
    ws.column_dimensions[L(mc + 1)].width = 17


def ex3_sheets(wb):
    seq_t, seq_s = read_csv("ex3_seq.csv"); omp_t, omp_s = read_csv("ex3_omp.csv")
    ws = wb.create_sheet("Ex3 Sequencial")
    ws["A1"] = "Exercício 3 – Multiplicação de matrizes (versão sequencial, blocking + i-k-j)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    header(ws, 3, ["Threads", "N", "Exec 1", "Exec 2", "Exec 3", "Média (s)"])
    ws.cell(row=4, column=1, value=1).font = INPUT
    ws.cell(row=4, column=2, value=seq_s.get(1)).font = INPUT
    for i in range(3):
        c = ws.cell(row=4, column=3 + i, value=seq_t[1][i] if len(seq_t.get(1, [])) > i else None)
        c.font = INPUT; c.number_format = "0.000000"
    ws["F4"] = '=IFERROR(AVERAGE(C4:E4),"")'; ws["F4"].number_format = "0.000000"
    for c in range(1, 7): ws.cell(row=4, column=c).border = BOX
    for c in "ABCDEF": ws.column_dimensions[c].width = 13

    wp = wb.create_sheet("Ex3 Paralela")
    wp["A1"] = "Exercício 3 – Multiplicação de matrizes (versão OpenMP, collapse(2) em blocos)"
    wp["A1"].font = Font(name=FONT, bold=True, size=14)
    wp["A2"] = "Speedup = média sequencial (aba 'Ex3 Sequencial') / média com p threads. Média de 3 execuções."
    wp["A2"].font = Font(name=FONT, italic=True, size=9)
    header(wp, 3, ["Threads", "N", "Exec 1", "Exec 2", "Exec 3", "Média (s)", "Speedup", "Eficiência", "Ideal"])
    for k, p in enumerate(THREADS):
        r = 4 + k
        wp.cell(row=r, column=1, value=p).font = INPUT
        wp.cell(row=r, column=2, value=omp_s.get(p)).font = INPUT
        for i in range(3):
            c = wp.cell(row=r, column=3 + i, value=omp_t[p][i] if len(omp_t.get(p, [])) > i else None)
            c.font = INPUT; c.number_format = "0.000000"
        wp.cell(row=r, column=6, value=f'=IFERROR(AVERAGE(C{r}:E{r}),"")').number_format = "0.000000"
        wp.cell(row=r, column=7, value=f"=IFERROR('Ex3 Sequencial'!$F$4/F{r},\"\")").number_format = "0.00"
        wp.cell(row=r, column=8, value=f'=IFERROR(G{r}/A{r},"")').number_format = "0.00"
        wp.cell(row=r, column=9, value=f"=A{r}")
        for c in range(1, 10): wp.cell(row=r, column=c).border = BOX
    for c in "ABCDEFGHI": wp.column_dimensions[c].width = 12
    ch = LineChart(); ch.title = "Speedup – multiplicação de matrizes"; ch.height = 8.5; ch.width = 16
    ch.x_axis.title = "Threads"; ch.y_axis.title = "Speedup"
    ch.x_axis.delete = False; ch.y_axis.delete = False
    for col in (7, 9):
        ch.add_data(Reference(wp, min_col=col, min_row=3, max_row=9), titles_from_data=True)
    ch.set_categories(Reference(wp, min_col=1, min_row=4, max_row=9))
    for s in ch.series: s.marker.symbol = "circle"
    wp.add_chart(ch, "K3")


wb = Workbook(); wb.remove(wb.active)
ex2_sheet(wb, "Ex2 Critical", "critical")
ex2_sheet(wb, "Ex2 Reduction", "reduction")
ex3_sheets(wb)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, size=c.font.size, color=c.font.color)
wb.save("resultados_openmp.xlsx")
print("Gerado: resultados_openmp.xlsx")
